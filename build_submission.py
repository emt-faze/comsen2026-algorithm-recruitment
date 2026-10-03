# -*- coding: utf-8 -*-
"""
把各 task 的 Markdown 报告转换成 Word 文档，并整理邮件提交材料。

运行方式（项目根目录）：
    python build_submission.py

生成内容：
    提交材料/task2报告.docx / task3报告.docx / task4报告.docx
    提交材料/task2/ 下的四个结果视频
"""
import os
import re
import shutil

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE_DIR, "提交材料")

REPORTS = [
    ("task2报告.docx", [os.path.join(BASE_DIR, "task2", "报告.md")]),
    ("task3报告.docx", [os.path.join(BASE_DIR, "task3", "报告.md"),
                        os.path.join(BASE_DIR, "task3", "q1_concepts.md")]),
    ("task4报告.docx", [os.path.join(BASE_DIR, "task4_timeseries", "报告.md")]),
]

VIDEOS = [
    ("task2/outputs/example2_frame_diff.mp4", "task2"),
    ("task2/outputs/example2_mog2.mp4", "task2"),
    ("task2/outputs/example3_tracked.mp4", "task2"),
    ("task2/outputs/example4_tracked.mp4", "task2"),
]


def set_base_style(doc):
    """统一正文字体：西文 Calibri + 中文微软雅黑。"""
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_runs_with_format(paragraph, text):
    """处理 **加粗** 和 `行内代码` 两种格式。"""
    # 先按加粗切分
    parts = re.split(r"\*\*(.+?)\*\*", text)
    for i, part in enumerate(parts):
        if not part:
            continue
        bold = i % 2 == 1
        # 加粗片段内部也可能有行内代码，简单起见直接原样输出
        run = paragraph.add_run(part)
        run.bold = bold
        if bold:
            run.font.color.rgb = None
    # 行内代码统一转成等宽字体颜色（不做精确解析）
    for run in paragraph.runs:
        if "`" in run.text:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_table(doc, rows):
    """rows: 去掉分隔行后的表格数据（第一行是表头）。"""
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            p = table.cell(r, c).paragraphs[0]
            run = p.add_run(cell)
            if r == 0:
                run.bold = True
    doc.add_paragraph()


def add_image(doc, rel_path, md_dir):
    img_path = os.path.normpath(os.path.join(md_dir, rel_path))
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(img_path, width=Inches(5.8))
    else:
        doc.add_paragraph(f"[缺少图片: {rel_path}]")


def parse_markdown(doc, md_path, skip_first_heading=False):
    md_dir = os.path.dirname(md_path)
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    in_code = False
    code_lines = []
    table_lines = []
    first_heading_seen = False

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        rows = []
        for line in table_lines:
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue  # 分隔行
            rows.append(cells)
        if rows:
            add_table(doc, rows)
        table_lines = []

    for line in lines:
        stripped = line.strip()

        # 代码块切换
        if stripped.startswith("```"):
            flush_table()
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue

        # 表格行
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines.append(stripped)
            continue
        flush_table()

        # 图片
        m = re.match(r"!\[.*?\]\((.+?)\)", stripped)
        if m:
            add_image(doc, m.group(1), md_dir)
            continue

        # 标题
        m = re.match(r"^(#{1,3})\s+(.*)", stripped)
        if m:
            if skip_first_heading and not first_heading_seen and m.group(1) == "#":
                first_heading_seen = True  # 附文件的标题不重复插入
                continue
            level = len(m.group(1))
            doc.add_heading("", level=level).add_run(m.group(2))
            continue

        # 列表
        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_format(p, re.sub(r"^[-*]\s+", "", stripped))
            continue
        m = re.match(r"^(\d+)[.、]\s+(.*)", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs_with_format(p, m.group(2))
            continue

        # 空行
        if not stripped:
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs_with_format(p, stripped)

    flush_table()


def build_docs():
    for out_name, md_files in REPORTS:
        doc = Document()
        set_base_style(doc)
        for i, md_path in enumerate(md_files):
            parse_markdown(doc, md_path, skip_first_heading=(i > 0))
        out_path = os.path.join(OUT_DIR, out_name)
        doc.save(out_path)
        print(f"已生成 {out_path}")


def copy_videos():
    for rel, sub in VIDEOS:
        src = os.path.join(BASE_DIR, rel)
        dst_dir = os.path.join(OUT_DIR, sub)
        os.makedirs(dst_dir, exist_ok=True)
        if os.path.exists(src):
            shutil.copy2(src, os.path.join(dst_dir, os.path.basename(src)))
            print(f"已复制 {rel}")
        else:
            print(f"[警告] 找不到 {src}")


if __name__ == "__main__":
    os.makedirs(OUT_DIR, exist_ok=True)
    build_docs()
    copy_videos()
    print("提交材料整理完成。")
